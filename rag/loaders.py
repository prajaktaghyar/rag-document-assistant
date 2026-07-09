"""
loaders.py
----------
Extracts plain text from many different file types so they can be chunked,
embedded, and indexed by the RAG pipeline.

Supported out of the box:
    .txt, .md, .csv, .tsv, .json, .log            (plain text family)
    .pdf                                           (pypdf)
    .docx                                          (python-docx)
    .pptx                                          (python-pptx)
    .xlsx, .xls                                    (pandas/openpyxl)
    .html, .htm                                    (BeautifulSoup)

Anything else falls back to a best-effort UTF-8 / latin-1 text decode, so
the app never hard-fails on an unrecognized extension -- it just does its
best to pull readable text out of the bytes.
"""

from __future__ import annotations

import io
import json
import os
from dataclasses import dataclass


@dataclass
class LoadedDocument:
    filename: str
    text: str
    num_chars: int
    file_type: str
    warning: str | None = None


def _decode_bytes_best_effort(data: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            return data.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    return data.decode("utf-8", errors="replace")


def _load_plain_text(data: bytes) -> str:
    return _decode_bytes_best_effort(data)


def _load_json(data: bytes) -> str:
    try:
        parsed = json.loads(data)
        return json.dumps(parsed, indent=2, ensure_ascii=False)
    except json.JSONDecodeError:
        return _decode_bytes_best_effort(data)


def _load_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages_text = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if text.strip():
            pages_text.append(f"[Page {i + 1}]\n{text}")
    return "\n\n".join(pages_text)


def _load_docx(data: bytes) -> str:
    import docx  # python-docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]

    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    return "\n".join(parts)


def _load_pptx(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    parts = []
    for slide_num, slide in enumerate(prs.slides, start=1):
        slide_parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text
                if text.strip():
                    slide_parts.append(text)
            if shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if any(cells):
                        slide_parts.append(" | ".join(cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            slide_parts.append(f"Notes: {slide.notes_slide.notes_text_frame.text}")
        if slide_parts:
            parts.append(f"[Slide {slide_num}]\n" + "\n".join(slide_parts))
    return "\n\n".join(parts)


def _load_excel(data: bytes, filename: str) -> str:
    import pandas as pd

    engine = "openpyxl" if filename.lower().endswith(("xlsx", "xlsm")) else None
    sheets = pd.read_excel(io.BytesIO(data), sheet_name=None, engine=engine)
    parts = []
    for sheet_name, df in sheets.items():
        parts.append(f"[Sheet: {sheet_name}]\n{df.to_csv(index=False)}")
    return "\n\n".join(parts)


def _load_csv(data: bytes, delimiter: str = ",") -> str:
    import pandas as pd

    df = pd.read_csv(io.BytesIO(data), delimiter=delimiter)
    return df.to_csv(index=False)


def _load_html(data: bytes) -> str:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(data, "lxml")
    for tag in soup(["script", "style"]):
        tag.decompose()
    text = soup.get_text(separator="\n")
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    return "\n".join(lines)


_EXTENSION_DISPATCH = {
    "txt": lambda data, name: _load_plain_text(data),
    "md": lambda data, name: _load_plain_text(data),
    "markdown": lambda data, name: _load_plain_text(data),
    "log": lambda data, name: _load_plain_text(data),
    "json": lambda data, name: _load_json(data),
    "csv": lambda data, name: _load_csv(data, delimiter=","),
    "tsv": lambda data, name: _load_csv(data, delimiter="\t"),
    "pdf": lambda data, name: _load_pdf(data),
    "docx": lambda data, name: _load_docx(data),
    "pptx": lambda data, name: _load_pptx(data),
    "xlsx": lambda data, name: _load_excel(data, name),
    "xls": lambda data, name: _load_excel(data, name),
    "html": lambda data, name: _load_html(data),
    "htm": lambda data, name: _load_html(data),
}


def load_document(filename: str, data: bytes) -> LoadedDocument:
    """
    Extract text from raw file bytes. Never raises for an unsupported
    extension -- falls back to best-effort decoding so the pipeline stays
    usable for "any" file type.
    """
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    warning = None

    handler = _EXTENSION_DISPATCH.get(ext)
    if handler is None:
        warning = (
            f"'.{ext or 'unknown'}' is not an explicitly supported format; "
            "attempted a raw text fallback, which may be incomplete or noisy."
        )
        text = _decode_bytes_best_effort(data)
    else:
        try:
            text = handler(data, filename)
        except Exception as exc:  # noqa: BLE001 - want to surface any parser error
            warning = f"Failed to parse with the '{ext}' handler ({exc}); used raw text fallback."
            text = _decode_bytes_best_effort(data)

    text = text.strip()
    return LoadedDocument(
        filename=filename,
        text=text,
        num_chars=len(text),
        file_type=ext or "unknown",
        warning=warning,
    )
